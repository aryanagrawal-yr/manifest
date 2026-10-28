"""Excel-backed storage.

Database file: data/manifest_db.xlsx
  Sheet "Dispatches": one row per dispatch (shared by manifest + gatepass)
  Sheet "Items":      one row per Order ID / AWB, linked by Dispatch ID
"""
import os
import re
import secrets
import shutil
import threading
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

IST = ZoneInfo("Asia/Kolkata")
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data"))
XLSX_PATH = os.path.join(DATA_DIR, "manifest_db.xlsx")
BACKUP_DIR = os.path.join(DATA_DIR, "backups")

LOCK = threading.RLock()

# (key, Excel header)
DISPATCH_COLS = [
    ("dispatch_id", "Dispatch ID"),
    ("manifest_no", "Manifest No"),
    ("gatepass_no", "Gatepass No"),
    ("dispatch_dt", "Dispatch Date & Time"),
    ("employee", "Employee Name"),
    ("partner", "Logistics Partner"),
    ("courier_name", "Courier Boy Name"),
    ("courier_phone", "Courier Boy Phone"),
    ("vehicle_no", "Vehicle Number"),
    ("vehicle_type", "Vehicle Type"),
    ("total_items", "Total Orders / AWB"),
    ("notes", "Notes"),
    ("entry_time", "Gate Entry Time"),
    ("exit_time", "Gate Exit Time"),
    ("created_by", "Created By"),
    ("created_at", "Created At"),
]
ITEM_COLS = [
    ("dispatch_id", "Dispatch ID"),
    ("sr", "Sr No"),
    ("ref", "Order ID / AWB"),
]
TEXT_KEYS = {"courier_phone", "vehicle_no", "dispatch_id"}


def now_ist():
    return datetime.now(IST)


def now_str():
    return now_ist().strftime("%Y-%m-%d %H:%M")


def secret_key():
    env = os.environ.get("SECRET_KEY")
    if env:
        return env
    os.makedirs(DATA_DIR, exist_ok=True)
    path = os.path.join(DATA_DIR, ".secret_key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass
    with open(path) as f:
        return f.read().strip()


def _style_header(ws, widths):
    fill = PatternFill("solid", fgColor="111111")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def _open():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(XLSX_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Dispatches"
        ws.append([h for _, h in DISPATCH_COLS])
        _style_header(ws, [20, 20, 20, 20, 18, 20, 20, 18, 16, 14, 12, 30, 18, 18, 14, 18])
        wi = wb.create_sheet("Items")
        wi.append([h for _, h in ITEM_COLS])
        _style_header(wi, [20, 8, 30])
        _save(wb)
    return load_workbook(XLSX_PATH)


def _save(wb):
    tmp = XLSX_PATH + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX_PATH)
    # one backup copy per day
    try:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        b = os.path.join(BACKUP_DIR, "manifest_db_%s.xlsx" % now_ist().strftime("%Y%m%d"))
        if not os.path.exists(b):
            shutil.copy2(XLSX_PATH, b)
    except OSError:
        pass


def _rows(ws, cols):
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        d = {}
        for i, (k, _) in enumerate(cols):
            v = row[i] if i < len(row) else None
            d[k] = "" if v is None else (str(v) if k != "total_items" and k != "sr" else v)
        out.append(d)
    return out


def list_dispatches():
    with LOCK:
        wb = _open()
        rows = _rows(wb["Dispatches"], DISPATCH_COLS)
        wb.close()
    return rows


def all_items():
    """dict: dispatch_id -> [refs]"""
    with LOCK:
        wb = _open()
        rows = _rows(wb["Items"], ITEM_COLS)
        wb.close()
    out = {}
    for r in rows:
        out.setdefault(r["dispatch_id"], []).append(r["ref"])
    return out


def get_dispatch(dispatch_id):
    with LOCK:
        wb = _open()
        disp = None
        for r in _rows(wb["Dispatches"], DISPATCH_COLS):
            if r["dispatch_id"] == dispatch_id:
                disp = r
                break
        if not disp:
            wb.close()
            return None
        disp["items"] = [
            r["ref"] for r in _rows(wb["Items"], ITEM_COLS) if r["dispatch_id"] == dispatch_id
        ]
        wb.close()
    return disp


def create_dispatch(f, items, created_by):
    """Returns (dispatch_id, duplicates) where duplicates are refs seen in earlier manifests."""
    with LOCK:
        wb = _open()
        ws = wb["Dispatches"]
        day = f["dispatch_dt"][:10].replace("-", "")
        prefix = "YR-%s-" % day
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and str(row[0]).startswith(prefix):
                try:
                    n = max(n, int(str(row[0]).rsplit("-", 1)[1]))
                except ValueError:
                    pass
        suffix = "%s-%03d" % (day, n + 1)
        did = "YR-" + suffix

        wi = wb["Items"]
        existing = {str(r[2]).strip().upper() for r in wi.iter_rows(min_row=2, values_only=True) if r[2]}
        dups = [x for x in items if x.upper() in existing]

        rec = dict(f)
        rec.update(
            dispatch_id=did,
            manifest_no="MF-" + suffix,
            gatepass_no="GP-" + suffix,
            total_items=len(items),
            entry_time="",
            exit_time="",
            created_by=created_by,
            created_at=now_str(),
        )
        ws.append([rec.get(k, "") for k, _ in DISPATCH_COLS])
        for i, ref in enumerate(items, start=1):
            wi.append([did, i, ref])
        # keep phone / vehicle / ids as text
        last = ws.max_row
        for idx, (k, _) in enumerate(DISPATCH_COLS, start=1):
            if k in TEXT_KEYS:
                ws.cell(row=last, column=idx).number_format = "@"
        _save(wb)
        wb.close()
    return did, dups


def mark_gate(dispatch_id, which):
    """which: 'entry' or 'exit'. Never overwrites a time that is already set."""
    col_key = "entry_time" if which == "entry" else "exit_time"
    col_idx = [k for k, _ in DISPATCH_COLS].index(col_key) + 1
    with LOCK:
        wb = _open()
        ws = wb["Dispatches"]
        done = False
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value) == dispatch_id:
                cell = row[col_idx - 1]
                if not cell.value:
                    cell.value = now_str()
                    done = True
                break
        if done:
            _save(wb)
        wb.close()
    return done


def split_refs(text):
    seen, out = set(), []
    for tok in re.split(r"[\s,;]+", text or ""):
        tok = tok.strip()
        if tok and tok.upper() not in seen:
            seen.add(tok.upper())
            out.append(tok)
    return out
